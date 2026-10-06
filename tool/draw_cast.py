"""Draws the cast as slide decks under art/, built from nothing but autoshapes.

This is how the first decks were made, not how art is maintained: once the
decks exist they are the source. Open them in any slide editor, move a shape,
and re-export with tool/export_art.py. Run this again only to start over --
it overwrites every deck it draws. Run from the repository root:

    python tool/draw_cast.py

Needs Python 3 and python-pptx.

Every deck follows docs/art-pipeline.md: one deck per actor and state
(art/<actor>_<state>.pptx), one pose per slide, and each pose a single group
on a 512 x 512 slide. Each group holds an invisible square, "Frame", covering
the whole slide, so the group's bounds -- which is what a slide editor's
"Save as Picture" exports -- are the same on every slide and the frames never
jitter.
"""

from __future__ import annotations

import math
import pathlib
from dataclasses import dataclass, field, replace

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parent.parent
ART = ROOT / "art"

CANVAS = 512
EMU_PER_PX = 9525  # 96 dpi

# The game's palette (lib/game/theme/palette.dart), plus a few shades.
INK = "1A2233"
SLIDE = "FFFCF5"
WHITE = "FFFFFF"
HIGHLIGHT = "F4B942"
HIGHLIGHT_DARK = "B7832A"
BRAND = "0E7C7B"
BRAND_DARK = "095C5B"
TIE = "D9534F"
TIE_DARK = "A33A37"
SHRINK = "6A4C93"
SHRINK_DARK = "4A3468"
DIAGRAM = "D1495B"
DIAGRAM_DARK = "A3364A"
HAT = "2A3F8F"
HAT_DARK = "1C2B63"
MASTER = "C77D1A"
MASTER_DARK = "8F5610"
SNAP = "D14FA8"
SNAP_DARK = "9C3480"
STEEL = "CFD3DC"
STEEL_DARK = "8A90A0"
NAVY = "1E2638"

OUTLINE = 12  # px at 512, so an outline survives being drawn at 50 px


# --- A tiny scene model --------------------------------------------------------


@dataclass
class Shape:
    kind: MSO_SHAPE
    cx: float
    cy: float
    w: float
    h: float
    fill: str | None = None
    line: str | None = None
    width: float = OUTLINE
    rot: float = 0
    alpha: float = 1
    adj: list[float] = field(default_factory=list)
    name: str = ""

    def moved(self, dx, dy):
        return replace(self, cx=self.cx + dx, cy=self.cy + dy)

    def turned(self, pivot, degrees):
        cx, cy = _turn((self.cx, self.cy), pivot, degrees)
        return replace(self, cx=cx, cy=cy, rot=self.rot + degrees)

    def stretched(self, pivot, sx, sy):
        px, py = pivot
        return replace(
            self,
            cx=px + (self.cx - px) * sx,
            cy=py + (self.cy - py) * sy,
            w=self.w * sx,
            h=self.h * sy,
            width=self.width * min(sx, sy) ** 0.5,
        )


@dataclass
class Line:
    x1: float
    y1: float
    x2: float
    y2: float
    colour: str = INK
    width: float = OUTLINE

    def moved(self, dx, dy):
        return replace(self, x1=self.x1 + dx, y1=self.y1 + dy, x2=self.x2 + dx, y2=self.y2 + dy)

    def turned(self, pivot, degrees):
        (x1, y1), (x2, y2) = _turn((self.x1, self.y1), pivot, degrees), _turn((self.x2, self.y2), pivot, degrees)
        return replace(self, x1=x1, y1=y1, x2=x2, y2=y2)

    def stretched(self, pivot, sx, sy):
        px, py = pivot
        return replace(
            self,
            x1=px + (self.x1 - px) * sx,
            y1=py + (self.y1 - py) * sy,
            x2=px + (self.x2 - px) * sx,
            y2=py + (self.y2 - py) * sy,
        )


def _turn(point, pivot, degrees):
    a = math.radians(degrees)
    x, y = point[0] - pivot[0], point[1] - pivot[1]
    return (pivot[0] + x * math.cos(a) - y * math.sin(a), pivot[1] + x * math.sin(a) + y * math.cos(a))


def moved(items, dx, dy):
    return [item.moved(dx, dy) for item in items]


def turned(items, pivot, degrees):
    return [item.turned(pivot, degrees) for item in items]


def stretched(items, pivot, sx, sy):
    return [item.stretched(pivot, sx, sy) for item in items]


def oval(cx, cy, w, h, fill=None, line=None, width=OUTLINE, **kw):
    return Shape(MSO_SHAPE.OVAL, cx, cy, w, h, fill, line, width, **kw)


def rounded(cx, cy, w, h, fill, line, radius=0.3, width=OUTLINE, **kw):
    return Shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, w, h, fill, line, width, adj=[radius], **kw)


def rect(cx, cy, w, h, fill, line=None, width=OUTLINE, **kw):
    return Shape(MSO_SHAPE.RECTANGLE, cx, cy, w, h, fill, line, width, **kw)


def triangle(cx, cy, w, h, fill, line=None, width=OUTLINE, **kw):
    return Shape(MSO_SHAPE.ISOSCELES_TRIANGLE, cx, cy, w, h, fill, line, width, **kw)


def arc(cx, cy, w, h, start, end, colour=INK, width=OUTLINE):
    """An unfilled arc, clockwise from start to end degrees (0 is east)."""
    return Shape(MSO_SHAPE.ARC, cx, cy, w, h, None, colour, width, adj=[start * 0.6, end * 0.6])


def shadow(cx, cy, w):
    return oval(cx, cy, w, w * 0.14, fill="000000", alpha=0.16)


# --- Faces ---------------------------------------------------------------------


def eyes(cx, cy, gap, r, mood):
    """Two eyes: open, blink, shut (it hurts), cross (beaten) or happy."""
    items = []
    for side in (-1, 1):
        ex = cx + side * gap
        if mood == "blink":
            items.append(Line(ex - r * 0.8, cy + r * 0.2, ex + r * 0.8, cy + r * 0.2))
            continue
        if mood == "shut":
            # Squeezed shut: a > and a <, pointing at the nose.
            tip = ex + side * -r * 0.5
            back = ex + side * r * 0.7
            items += [Line(back, cy - r * 0.6, tip, cy), Line(tip, cy, back, cy + r * 0.6)]
            continue
        items.append(oval(ex, cy, r * 2, r * 2, fill=WHITE, line=INK, width=OUTLINE * 0.75))
        if mood == "open":
            items.append(oval(ex + r * 0.18, cy + r * 0.12, r * 0.9, r * 0.9, fill=INK))
        elif mood == "cross":
            d = r * 0.55
            items += [Line(ex - d, cy - d, ex + d, cy + d, width=OUTLINE * 0.8), Line(ex - d, cy + d, ex + d, cy - d, width=OUTLINE * 0.8)]
        elif mood == "happy":
            items.append(arc(ex, cy + r * 0.35, r * 1.2, r * 1.0, 180, 0, width=OUTLINE * 0.8))
    return items


def mouth(cx, cy, w, kind, colour=INK):
    if kind == "flat":
        return [Line(cx - w / 2, cy, cx + w / 2, cy, colour)]
    if kind == "o":
        return [oval(cx, cy, w * 0.32, w * 0.32, fill=INK, line=colour, width=OUTLINE * 0.6)]
    if kind == "grin":
        return [arc(cx, cy - w * 0.25, w, w * 0.6, 0, 180, colour)]
    return []


def expression(state, frame, blink_frame=13):
    """The eyes and mouth each state wears on each frame."""
    if state == "hit":
        return "shut", "o"
    if state == "die":
        return "cross", "o"
    if state == "cheer":
        return "happy", "grin"
    if state == "idle" and frame == blink_frame:
        return "blink", "flat"
    return "open", "flat"


# --- The cast ------------------------------------------------------------------

FEET = (256, 470)


def hit_pose(items, frame, keep):
    """A squash towards the feet, springing back on the second frame."""
    sx, sy = [(1.15, 0.82), (1.06, 0.93)][frame]
    return keep + stretched(items, FEET, sx, sy)


def die_pose(items, frame, keep):
    """Toppling: upright, tipping, and down."""
    degrees, sy = [(0, 1), (12, 0.97), (24, 0.88)][frame]
    return keep + turned(stretched(items, FEET, 1, sy), FEET, degrees)


def hero(state, frame):
    """The presenter: a teal block with a bullet point for an antenna, and a tie."""
    look, say = expression(state, frame)
    legs = [rounded(210, 440, 48, 62, BRAND_DARK, BRAND_DARK, 0.4), rounded(302, 440, 48, 62, BRAND_DARK, BRAND_DARK, 0.4)]
    arms = [
        oval(126, 318, 46, 96, fill=BRAND, line=BRAND_DARK, rot=12),
        oval(386, 318, 46, 96, fill=BRAND, line=BRAND_DARK, rot=-12),
    ]
    if state == "walk":
        step = [1, 0, -1, 0][frame]
        legs = [
            rounded(210, 440 - 12 * max(step, 0), 48, 62 - 16 * max(step, 0), BRAND_DARK, BRAND_DARK, 0.4),
            rounded(302, 440 - 12 * max(-step, 0), 48, 62 - 16 * max(-step, 0), BRAND_DARK, BRAND_DARK, 0.4),
        ]
        arms = [arms[0].turned((126, 280), -22 * step), arms[1].turned((386, 280), -22 * step)]
    if state == "cheer":
        lift = [0, 14][frame]
        arms = [
            oval(124, 196 - lift, 46, 104, fill=BRAND, line=BRAND_DARK, rot=-28),
            oval(388, 196 - lift, 46, 104, fill=BRAND, line=BRAND_DARK, rot=28),
        ]
    body = [
        Line(256, 192, 256, 126, BRAND_DARK),
        oval(256, 108, 46, 46, fill=BRAND, line=BRAND_DARK),
        rounded(256, 300, 250, 236, BRAND, BRAND_DARK, 0.32),
        rounded(256, 236, 222, 74, "2A9493", None, 0.45, width=0),
        rect(256, 378, 34, 20, TIE, TIE_DARK, width=OUTLINE * 0.6),
        triangle(256, 412, 46, 50, TIE, TIE_DARK, width=OUTLINE * 0.6, rot=180),
        *eyes(256, 278, 54, 34, look),
        *mouth(256, 346, 72, say),
    ]
    upper = arms + body
    if state == "walk":
        upper = turned(upper, FEET, [4, 0, -4, 0][frame])
    if state == "cheer":
        upper = moved(upper, 0, -[0, 8][frame])
    pose = legs + upper
    floor = [shadow(256, 478, 300)]
    if state == "hit":
        return hit_pose(pose, frame, floor)
    if state == "die":
        return die_pose(pose, frame, floor)
    return floor + pose


def shrink_to_fit(state, frame):
    """A purple text box with an "A" for a belly, its corners squeezing in."""
    look, _ = expression(state, frame)
    squeeze = [0, 4, 8, 12, 8, 4, 0, 0][frame % 8] if state == "idle" else 18
    arrows = []
    for sx, sy, rot in [(-1, -1, 45), (1, -1, 135), (1, 1, 225), (-1, 1, 315)]:
        d = 172 - squeeze
        arrows.append(
            Shape(MSO_SHAPE.RIGHT_ARROW, 256 + sx * d, 290 + sy * d * 0.86, 78, 48, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.6, rot=rot)
        )
    pose = [
        oval(212, 432, 64, 30, fill=SHRINK_DARK),
        oval(300, 432, 64, 30, fill=SHRINK_DARK),
        rounded(256, 290, 262, 250, SHRINK, SHRINK_DARK, 0.18),
        triangle(256, 360, 104, 92, SLIDE, None, width=0),
        triangle(256, 376, 40, 36, SHRINK, None, width=0),
        rect(256, 384, 70, 12, SHRINK, None, width=0),
        *eyes(256, 244, 54, 32, look),
        *arrows,
    ]
    floor = [shadow(256, 466, 290)]
    if state == "hit":
        return hit_pose(pose, frame, floor)
    if state == "die":
        # Shrink-to-Fit's own medicine.
        s = [0.85, 0.65, 0.45][frame]
        return floor + stretched(pose, FEET, s, s)
    return floor + pose


def diagram_wizard(state, frame):
    """A diagram shape in a wizard's hat."""
    look, say = expression(state, frame)
    twinkle = [0, 12, 24, 36, 24, 12, 0, 0][frame % 8] if state == "idle" else 0
    hat = [
        triangle(256, 150, 226, 176, HAT, HAT_DARK),
        rect(256, 236, 300, 28, HAT, HAT_DARK, width=OUTLINE * 0.75),
        Shape(MSO_SHAPE.STAR_5_POINT, 238, 168, 56, 56, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.5, rot=twinkle),
    ]
    if state == "die":
        centre = (256, 200)
        hat = moved(turned(hat, centre, [0, 22, 48][frame]), [0, 40, 110][frame], [0, -10, 60][frame])
    body = [
        rounded(256, 336, 284, 224, DIAGRAM, DIAGRAM_DARK, 0.26),
        *eyes(256, 320, 58, 34, look),
        *mouth(256, 388, 70, say),
    ]
    floor = [shadow(256, 470, 300)]
    if state == "hit":
        return hit_pose(body + hat, frame, floor)
    if state == "die":
        return die_pose(body, frame, floor) + hat
    return floor + body + hat


def master_template(state, frame):
    """The master: an amber slide with a crown and a monocle."""
    look, _ = expression(state, frame)
    crown = [
        triangle(186, 168, 64, 76, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.6),
        triangle(256, 156, 70, 100, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.6),
        triangle(326, 168, 64, 76, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.6),
        rect(256, 214, 200, 40, HIGHLIGHT, HIGHLIGHT_DARK, OUTLINE * 0.6),
        oval(256, 214, 22, 22, fill=TIE),
    ]
    if state == "die":
        crown = moved(turned(crown, (256, 200), [0, 25, 55][frame]), [0, 50, 120][frame], [0, -6, 70][frame])
    glint = state == "idle" and frame in (2, 3)
    face = [
        rounded(256, 322, 312, 236, MASTER, MASTER_DARK, 0.2),
        rect(256, 236, 260, 22, "E09A3A", None, width=0),
        *eyes(256, 306, 62, 34, look),
        oval(318, 306, 98, 98, fill=None, line=HIGHLIGHT_DARK, width=OUTLINE * 0.75),
        Line(352, 342, 378, 410, HIGHLIGHT_DARK, OUTLINE * 0.5),
        rounded(256, 378, 104, 22, MASTER_DARK, MASTER_DARK, 0.5),
    ]
    if glint:
        face.append(Line(296, 280, 310, 266, WHITE, OUTLINE * 0.6))
    floor = [shadow(256, 466, 300)]
    if state == "hit":
        return hit_pose(face + crown, frame, floor)
    if state == "die":
        return die_pose(face, frame, floor) + crown
    return floor + face + crown


def snap_to_grid(state, frame):
    """A horseshoe magnet: everything snaps to it."""
    look, _ = expression(state, frame)
    pose = [
        Shape(MSO_SHAPE.BLOCK_ARC, 256, 310, 340, 340, SNAP, SNAP_DARK, OUTLINE, adj=[180 * 0.6, 0, 0.28]),
        rect(133, 344, 96, 66, STEEL, STEEL_DARK, OUTLINE * 0.75),
        rect(379, 344, 96, 66, STEEL, STEEL_DARK, OUTLINE * 0.75),
        *eyes(256, 188, 58, 30, look),
    ]
    if state == "idle":
        # Little sparks off the poles, alternating.
        if frame % 4 in (1, 2):
            pose += [Line(110, 400, 96, 426, HIGHLIGHT, OUTLINE * 0.6), Line(156, 400, 170, 426, HIGHLIGHT, OUTLINE * 0.6)]
        if frame % 4 in (3, 0) and frame != 0:
            pose += [Line(356, 400, 342, 426, HIGHLIGHT, OUTLINE * 0.6), Line(402, 400, 416, 426, HIGHLIGHT, OUTLINE * 0.6)]
    floor = [shadow(256, 450, 280)]
    if state == "hit":
        return hit_pose(pose, frame, floor)
    if state == "die":
        return die_pose(pose, frame, floor)
    return floor + pose


# Idle runs 16 frames at the game's 0.14 s a frame, so the blink comes
# every two seconds or so rather than every one.
FRAMES = {"idle": 16, "walk": 4, "hit": 2, "die": 3, "cheer": 2}

CAST = {
    "hero": (hero, ["idle", "walk", "hit", "die", "cheer"]),
    "shrink_to_fit": (shrink_to_fit, ["idle", "hit", "die"]),
    "diagram_wizard": (diagram_wizard, ["idle", "hit", "die"]),
    "master_template": (master_template, ["idle", "hit", "die"]),
    "snap_to_grid": (snap_to_grid, ["idle", "hit", "die"]),
}


def launcher_icon():
    """The app icon: the presenter on the editor's navy, full bleed."""
    return [rect(256, 256, 512, 512, NAVY)] + stretched(moved(hero("idle", 0), 0, -22), (256, 256), 0.86, 0.86)


def launcher_foreground():
    """The adaptive icon's foreground layer: the presenter alone, small enough
    to stay inside the safe zone -- the middle 66 of 108 -- that every
    launcher's mask leaves alone. The navy is the background layer."""
    return stretched(moved(hero("idle", 0), 0, -22), (256, 256), 0.56, 0.56)


# --- Writing decks -------------------------------------------------------------


def px(value):
    return Emu(round(value * EMU_PER_PX))


def _fill(fill, hex_, alpha=1):
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_)
    if alpha < 1:
        colour = fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(colour, qn("a:alpha"), val=str(round(alpha * 100000)))


def _add(group, item):
    if isinstance(item, Line):
        line = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(item.x1), px(item.y1), px(item.x2), px(item.y2))
        line.line.color.rgb = RGBColor.from_string(item.colour)
        line.line.width = px(item.width)
        line.line._get_or_add_ln().set("cap", "rnd")
        return
    shape = group.shapes.add_shape(item.kind, px(item.cx - item.w / 2), px(item.cy - item.h / 2), px(item.w), px(item.h))
    for i, value in enumerate(item.adj):
        shape.adjustments[i] = value
    if item.fill:
        _fill(shape.fill, item.fill, item.alpha)
    else:
        shape.fill.background()
    if item.line and item.width > 0:
        shape.line.color.rgb = RGBColor.from_string(item.line)
        shape.line.width = px(item.width)
    else:
        shape.line.fill.background()
    shape.rotation = item.rot % 360
    shape.shadow.inherit = False
    if item.name:
        shape.name = item.name


def write_deck(path, poses, name):
    prs = Presentation()
    prs.slide_width = prs.slide_height = px(CANVAS)
    blank = prs.slide_layouts[6]
    for i, items in enumerate(poses):
        slide = prs.slides.add_slide(blank)
        group = slide.shapes.add_group_shape()
        group.name = f"{name} {i:03d}"
        frame = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(CANVAS), px(CANVAS))
        frame.name = "Frame"
        frame.fill.background()
        frame.line.fill.background()
        frame.shadow.inherit = False
        for item in items:
            _add(group, item)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    print(f"{path.relative_to(ROOT)}  {len(poses)} slides")


def main() -> None:
    for actor, (draw, states) in CAST.items():
        for state in states:
            poses = [draw(state, frame) for frame in range(FRAMES[state])]
            write_deck(ART / f"{actor}_{state}.pptx", poses, f"{actor} {state}")
    write_deck(ART / "launcher_icon.pptx", [launcher_icon(), launcher_foreground()], "launcher icon")


if __name__ == "__main__":
    main()
