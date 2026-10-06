"""Exports every deck under art/ to PNG: the one command that regenerates the
game's artwork. Run from the repository root:

    python tool/export_art.py            every deck
    python tool/export_art.py art/hero_walk.pptx

Needs Python 3 with python-pptx, numpy and Pillow, LibreOffice (soffice) and
poppler's pdftocairo -- all free, and the same on Windows, macOS and Linux.
With PowerPoint on Windows, tool/export-art.ps1 does the actor decks through
PowerPoint itself instead.

* art/<actor>_<state>.pptx becomes assets/images/<actor>_<state>_NNN.png,
  one 512 x 512 frame per slide on a transparent background. A sequence's old
  frames are deleted first, so a removed slide leaves no stale frame behind.
* art/launcher_icon.pptx becomes every launcher icon the platforms ask for:
  Android's mipmaps, the iOS app icon set and the Windows .ico.

The drawing rules in docs/art-pipeline.md are checked before anything is
written: a slide that is not exactly one grouped shape fails the run.

LibreOffice paints every page white when it exports, so each deck is rendered
twice, once on white and once on black. Where the two agree the pixel is
solid; how far they differ is how transparent it is. That recovers the exact
colour and coverage of every pixel, soft edges included.
"""

from __future__ import annotations

import argparse
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = pathlib.Path(__file__).resolve().parent.parent
ART = ROOT / "art"
IMAGES = ROOT / "assets" / "images"

FRAME_SIZE = 512
ICON_SIZE = 1024
DECK_NAME = re.compile(r"^[a-z0-9_]+_(idle|walk|hit|die|cheer)(_(s|se|e|ne|n))?$")
ICON_DECK = "launcher_icon"


class ArtError(Exception):
    """A deck that breaks the drawing rules."""


def check(deck: pathlib.Path) -> int:
    """Holds a deck to the drawing rules and returns its slide count."""
    prs = Presentation(deck)
    if prs.slide_width != prs.slide_height:
        raise ArtError(f"{deck.name}: slides must be square, to export without stretching")
    for number, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        if len(shapes) != 1:
            raise ArtError(
                f"{deck.name}, slide {number}: {len(shapes)} top-level shapes. "
                "Group each pose into one shape (select all, Ctrl+G)."
            )
        if shapes[0].shape_type != 6:  # MSO_SHAPE_TYPE.GROUP
            raise ArtError(f"{deck.name}, slide {number}: the pose is not a group.")
    if len(prs.slides) == 0:
        raise ArtError(f"{deck.name}: no slides")
    return len(prs.slides)


def _on_background(deck: pathlib.Path, out: pathlib.Path, colour: str) -> None:
    """A copy of [deck] with every slide on a solid [colour] background."""
    prs = Presentation(deck)
    for slide in prs.slides:
        c_sld = slide._element.find(qn("p:cSld"))
        old = c_sld.find(qn("p:bg"))
        if old is not None:
            c_sld.remove(old)
        bg = etree.Element(qn("p:bg"))
        c_sld.insert(0, bg)
        bg_pr = etree.SubElement(bg, qn("p:bgPr"))
        fill = etree.SubElement(bg_pr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=colour)
        etree.SubElement(bg_pr, qn("a:effectLst"))
    prs.save(out)


def _render(deck: pathlib.Path, work: pathlib.Path, size: int) -> list[Image.Image]:
    """Renders every slide of [deck] at [size] px, transparent where empty."""
    pages = {}
    profile = (work / "profile").as_uri()
    for name, colour in [("white", "FFFFFF"), ("black", "000000")]:
        copy = work / f"{deck.stem}_{name}.pptx"
        _on_background(deck, copy, colour)
        subprocess.run(
            [
                _tool("soffice"),
                f"-env:UserInstallation={profile}",
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(work),
                str(copy),
            ],
            check=True,
            capture_output=True,
        )
        prefix = work / f"{deck.stem}_{name}"
        subprocess.run(
            [_tool("pdftocairo"), "-png", "-scale-to", str(size), str(prefix.with_suffix(".pdf")), str(prefix)],
            check=True,
            capture_output=True,
        )
        pages[name] = sorted(work.glob(f"{deck.stem}_{name}-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
    return [_matte(white, black) for white, black in zip(pages["white"], pages["black"])]


def _matte(on_white: pathlib.Path, on_black: pathlib.Path) -> Image.Image:
    white = np.asarray(Image.open(on_white).convert("RGB"), dtype=float)
    black = np.asarray(Image.open(on_black).convert("RGB"), dtype=float)
    alpha = (1 - (white - black).mean(axis=2) / 255).clip(0, 1)
    colour = np.where(alpha[..., None] > 0, black / np.maximum(alpha[..., None], 1e-6), 0)
    rgba = np.dstack([colour.clip(0, 255), alpha * 255]).round().astype(np.uint8)
    return Image.fromarray(rgba, "RGBA")


def _tool(name: str) -> str:
    found = shutil.which(name)
    if found is None:
        raise ArtError(f"{name} is not installed (see the top of tool/export_art.py)")
    return found


def export_frames(deck: pathlib.Path, work: pathlib.Path) -> None:
    prefix = f"{deck.stem}_"
    frames = _render(deck, work, FRAME_SIZE)
    stale = [p for p in IMAGES.glob(f"{prefix}*.png") if re.fullmatch(re.escape(prefix) + r"\d{3}\.png", p.name)]
    for path in stale:
        path.unlink()
    IMAGES.mkdir(parents=True, exist_ok=True)
    for i, frame in enumerate(frames):
        frame.save(IMAGES / f"{prefix}{i:03d}.png", optimize=True)
    print(f"{deck.name:28} {len(frames)} frames")


def export_icon(deck: pathlib.Path, work: pathlib.Path) -> None:
    # App icons are opaque: iOS refuses one with an alpha channel.
    icon = _render(deck, work, ICON_SIZE)[0]
    flat = Image.new("RGB", icon.size, "white")
    flat.paste(icon, mask=icon.getchannel("A"))
    written = []

    def save(path: pathlib.Path, size: int) -> None:
        flat.resize((size, size), Image.LANCZOS).save(path, optimize=True)
        written.append(path)

    res = ROOT / "android" / "app" / "src" / "main" / "res"
    for density, size in [("mdpi", 48), ("hdpi", 72), ("xhdpi", 96), ("xxhdpi", 144), ("xxxhdpi", 192)]:
        save(res / f"mipmap-{density}" / "ic_launcher.png", size)
    icon_set = ROOT / "ios" / "Runner" / "Assets.xcassets" / "AppIcon.appiconset"
    for path in sorted(icon_set.glob("Icon-App-*.png")):
        match = re.fullmatch(r"Icon-App-([\d.]+)x[\d.]+@(\d)x\.png", path.name)
        if match:
            save(path, round(float(match[1]) * int(match[2])))
    ico = ROOT / "windows" / "runner" / "resources" / "app_icon.ico"
    flat.save(ico, sizes=[(s, s) for s in (16, 24, 32, 48, 64, 128, 256)])
    written.append(ico)
    print(f"{deck.name:28} {len(written)} icons")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("decks", nargs="*", type=pathlib.Path, help="decks to export (default: every deck in art/)")
    decks = [d.resolve() for d in parser.parse_args().decks] or sorted(ART.glob("*.pptx"))
    try:
        # Every deck is checked before any file is touched.
        for deck in decks:
            if deck.stem != ICON_DECK and not DECK_NAME.fullmatch(deck.stem):
                raise ArtError(f"{deck.name}: expected <actor>_<state>.pptx (see docs/art-pipeline.md)")
            check(deck)
        with tempfile.TemporaryDirectory() as tmp:
            for deck in decks:
                work = pathlib.Path(tmp) / deck.stem
                work.mkdir()
                if deck.stem == ICON_DECK:
                    export_icon(deck, work)
                else:
                    export_frames(deck, work)
    except ArtError as error:
        print(f"export_art: {error}", file=sys.stderr)
        return 1
    except subprocess.CalledProcessError as error:
        print(f"export_art: {error.cmd[0]} failed:\n{error.stderr.decode(errors='replace')}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
